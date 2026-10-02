import re

path = 'src/contabilidad/vista_reportes.py'
with open(path, 'r', encoding='utf-8') as f:
    content = f.read()

# Reemplazar botones CSV por EXCEL
content = content.replace('"📊  Exportar Gastos CSV"', '"📊  Exportar Gastos Excel"')
content = content.replace('"📊  Exportar Ingresos CSV"', '"📊  Exportar Ingresos Excel"')
content = content.replace('"📈  Exportar Préstamos CSV"', '"📈  Exportar Préstamos Excel"')

content = content.replace('self._export_gastos_csv', 'self._export_gastos_excel')
content = content.replace('self._export_ingresos_csv', 'self._export_ingresos_excel')
content = content.replace('self._export_prestamos_csv', 'self._export_prestamos_excel')

# Buscar el método de exportar csv e inyectar el de excel
excel_methods = """
    def _export_excel_generic(self, rows, headers, sheet_title, filename):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        path, _ = QFileDialog.getSaveFileName(self, "Exportar Excel", filename, "Excel (*.xlsx)")
        if not path: return
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_title

        # Estilos de encabezado
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="4F46E5") # Color primario TPV PRO
        align_center = Alignment(horizontal="center", vertical="center")

        # Escribir encabezados
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center

        # Escribir filas
        for row_num, row_data in enumerate(rows or [], 2):
            for col_num, cell_value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_num, value=cell_value)
                # Si el encabezado dice "Monto", aplicamos formato de moneda
                if "Monto" in headers[col_num - 1]:
                    try:
                        cell.value = float(cell_value)
                        cell.number_format = '$#,##0.00'
                    except:
                        pass
        
        # Ajustar ancho de columnas
        for col_num, column_cells in enumerate(ws.columns, 1):
            max_length = 0
            column_letter = get_column_letter(col_num)
            for cell in column_cells:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column_letter].width = adjusted_width

        wb.save(path)
        QMessageBox.information(self.window(), "Exportado", f"Guardado en:\\n{path}")

    def _export_gastos_excel(self):
        try:
            self._export_excel_generic(self._db.get_expenses(),
                                      ["ID", "Fecha", "Categoría", "Monto", "Descripción", "Tipo"],
                                      "Gastos",
                                      f"gastos_{self._mes:02d}_{self._año}.xlsx")
        except Exception as e: QMessageBox.critical(self.window(), "Error", str(e))

    def _export_ingresos_excel(self):
        try:
            self._export_excel_generic(self._db.get_income(self._mes, self._año),
                                      ["ID", "Fecha", "Descripción", "Fuente", "Monto"],
                                      "Ingresos",
                                      f"ingresos_{self._mes:02d}_{self._año}.xlsx")
        except Exception as e: QMessageBox.critical(self.window(), "Error", str(e))

    def _export_prestamos_excel(self):
        try:
            self._export_excel_generic(self._db.get_installments(),
                                      ["ID", "LoanID", "N°", "Monto", "Vencimiento", "Estado", "PaidDate"],
                                      "Préstamos",
                                      f"prestamos_{self._año}.xlsx")
        except Exception as e: QMessageBox.critical(self.window(), "Error", str(e))

"""

# Reemplazar los metodos csv viejos por los de excel (o añadirlos arriba de _do_backup)
import re
content = re.sub(r'    def _export_csv_generic\(self.*?    # BACKUP', excel_methods + r'    # BACKUP', content, flags=re.DOTALL)

with open(path, 'w', encoding='utf-8') as f:
    f.write(content)

print("Patch applied")
