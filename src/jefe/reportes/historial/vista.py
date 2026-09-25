"""Cara clara del historial. El motor firma; acá solo se lee."""

from datetime import datetime

from PyQt6.QtCore import QDate, QTime, Qt
from PyQt6.QtGui import QBrush, QColor
from PyQt6.QtWidgets import (
    QAbstractItemView,
    QCheckBox,
    QComboBox,
    QDateEdit,
    QDialog,
    QFileDialog,
    QFrame,
    QGridLayout,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QTimeEdit,
    QVBoxLayout,
    QWidget,
)

from src.historial_ventas import motor_historial
from src.jefe.reportes.financiero.dinero import fmt_plata
from src.jefe.reportes.letra import etiqueta, fuente_limpia, paleta_clara, vestir_fecha
from src.jefe.reportes.vista_financiero import _FIN, _aplicar_paleta_tabla
from src.utils.qt_compat import qt_exec


def _lbl(texto: str, px: int = 13, _peso: int = 400, _color: str | None = None) -> QLabel:
    return etiqueta(texto, px)


class VistaHistorial(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setObjectName("JefeHistorial")
        self.ticket_id = None
        self._armado = False
        self._setup()
        self.refrescar()

    def _setup(self):
        paleta_clara(self)
        self.setFont(fuente_limpia(13))
        self.setStyleSheet(
            f"""
            QWidget#JefeHistorial {{
                background: {_FIN['bg_page']};
                letter-spacing: 0px;
                font-family: 'Segoe UI';
                font-weight: 400 !important;
            }}
            QWidget#JefeHistorial QLabel,
            QWidget#JefeHistorial QPushButton,
            QWidget#JefeHistorial QCheckBox,
            QWidget#JefeHistorial QComboBox,
            QWidget#JefeHistorial QLineEdit,
            QWidget#JefeHistorial QDateEdit,
            QWidget#JefeHistorial QTimeEdit,
            QWidget#JefeHistorial QHeaderView::section {{
                letter-spacing: 0px;
                font-weight: 400 !important;
            }}
            QWidget#JefeHistorial QLineEdit,
            QWidget#JefeHistorial QComboBox,
            QWidget#JefeHistorial QDateEdit,
            QWidget#JefeHistorial QTimeEdit {{
                background: {_FIN['card']};
                border: 1px solid {_FIN['card_border']};
                border-radius: 8px;
                padding: 6px 10px;
                min-height: 28px;
                font-weight: 400 !important;
            }}
            QWidget#JefeHistorial QPushButton {{
                background: {_FIN['card']};
                border: 1px solid {_FIN['card_border']};
                border-radius: 8px;
                padding: 8px 14px;
                font-weight: 400 !important;
            }}
            QWidget#JefeHistorial QPushButton:hover {{
                background: {_FIN['accent_light']};
            }}
            QWidget#JefeHistorial QPushButton:disabled {{
                background: {_FIN['row_gray']};
            }}
            QWidget#JefeHistorial QCheckBox {{
                background: transparent;
                font-weight: 400 !important;
                letter-spacing: 0px;
            }}
            """
        )
        root = QVBoxLayout(self)
        root.setContentsMargins(24, 18, 24, 18)
        root.setSpacing(10)

        root.addWidget(_lbl("Historial de ventas", 22, 600))
        root.addWidget(
            _lbl(
                "Misma caja que el terminal. Acá se ve quién canceló y desde qué caja. "
                "La ganancia se firma en Reportes solo si hay costo cargado.",
                12,
                400,
                _FIN["text_soft"],
            )
        )

        cuerpo = QHBoxLayout()
        cuerpo.setSpacing(16)

        izq = QVBoxLayout()
        izq.addWidget(_lbl("Buscar por folio, producto o cajero", 12, 400, _FIN["text_soft"]))
        self.txt = QLineEdit()
        self.txt.setPlaceholderText("Ej: 1024, asado o cajero")
        self.txt.textChanged.connect(self.refrescar)
        izq.addWidget(self.txt)

        self.tabla = QTableWidget()
        self.tabla.setColumnCount(6)
        self.tabla.setHorizontalHeaderLabels(["Folio", "Caja", "Arts", "Hora", "Total", "Estado"])
        self.tabla.verticalHeader().setVisible(False)
        self.tabla.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.tabla.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.tabla.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.tabla.horizontalHeader().setSectionResizeMode(4, QHeaderView.ResizeMode.Stretch)
        self.tabla.itemSelectionChanged.connect(self._detalle)
        _aplicar_paleta_tabla(self.tabla)
        izq.addWidget(self.tabla)

        filt = QGridLayout()
        filt.setHorizontalSpacing(8)
        filt.setVerticalSpacing(8)
        filt.addWidget(_lbl("Día"), 0, 0)
        self.fecha = QDateEdit(QDate.currentDate())
        vestir_fecha(self.fecha)
        self._pintar_dias_calendario()
        self.fecha.dateChanged.connect(self._al_cambiar_dia)
        filt.addWidget(self.fecha, 0, 1)
        hoy = QPushButton("Hoy")
        hoy.clicked.connect(self._ir_hoy)
        filt.addWidget(hoy, 0, 2)
        rec = QPushButton("Actualizar")
        rec.clicked.connect(self.refrescar)
        filt.addWidget(rec, 0, 3)

        self.chk_todo = QCheckBox("Todas las fechas (histórico)")
        self.chk_todo.setChecked(False)
        self.chk_todo.setStyleSheet("background: transparent; letter-spacing: 0px; font-weight: 400;")
        self.chk_todo.stateChanged.connect(self.refrescar)
        filt.addWidget(self.chk_todo, 1, 0, 1, 4)

        filt.addWidget(_lbl("Pago"), 2, 0)
        self.cb_pago = QComboBox()
        self.cb_pago.addItems(
            ["TODOS", "EFECTIVO", "TARJETA", "TRANSFERENCIA", "MIXTO", "FIADO", "CLIENTES"]
        )
        self.cb_pago.currentIndexChanged.connect(self.refrescar)
        filt.addWidget(self.cb_pago, 2, 1)

        filt.addWidget(_lbl("Caja"), 2, 2)
        self.cb_caja = QComboBox()
        self.cb_caja.addItem("TODAS", None)
        for n in motor_historial.cajas():
            self.cb_caja.addItem(f"Caja {n}", n)
        self.cb_caja.currentIndexChanged.connect(self.refrescar)
        filt.addWidget(self.cb_caja, 2, 3)

        self.t0 = QTimeEdit(QTime(0, 0))
        self.t1 = QTimeEdit(QTime(23, 59))
        self.t0.setDisplayFormat("HH:mm")
        self.t1.setDisplayFormat("HH:mm")
        self.t0.timeChanged.connect(self.refrescar)
        self.t1.timeChanged.connect(self.refrescar)
        horas = QHBoxLayout()
        horas.addWidget(_lbl("De"))
        horas.addWidget(self.t0)
        horas.addWidget(_lbl("a"))
        horas.addWidget(self.t1)
        filt.addLayout(horas, 3, 1, 1, 3)
        izq.addLayout(filt)
        cuerpo.addLayout(izq, 48)

        der = QVBoxLayout()
        self.lbl_ticket = _lbl("Elegí un ticket", 18, 600)
        der.addWidget(self.lbl_ticket)

        card = QFrame()
        card.setObjectName("ticketCard")
        card.setStyleSheet(
            f"#ticketCard {{ background: {_FIN['card']}; border: 1px solid {_FIN['card_border']}; "
            f"border-radius: 16px; }}"
        )
        cl = QVBoxLayout(card)
        cl.setSpacing(6)
        self.lbl_cajero = _lbl("Cajero  —")
        self.lbl_pago = _lbl("Pago  —")
        self.lbl_caja = _lbl("Caja  —")
        self.lbl_cuando = _lbl("")
        self.lbl_audit = _lbl("", 12, 400, "#B45309")
        self.lbl_audit.setWordWrap(True)
        self.lbl_audit.hide()
        cl.addWidget(self.lbl_cajero)
        cl.addWidget(self.lbl_pago)
        cl.addWidget(self.lbl_caja)
        cl.addWidget(self.lbl_cuando)
        cl.addWidget(self.lbl_audit)
        self.tabla_det = QTableWidget()
        self.tabla_det.setColumnCount(3)
        self.tabla_det.setHorizontalHeaderLabels(["Cant.", "Producto", "Importe"])
        self.tabla_det.verticalHeader().setVisible(False)
        self.tabla_det.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        _aplicar_paleta_tabla(self.tabla_det)
        cl.addWidget(self.tabla_det)
        self.lbl_total = _lbl("Total  —", 18, 600)
        cl.addWidget(self.lbl_total)
        der.addWidget(card)

        fila = QHBoxLayout()
        self.btn_cancel = QPushButton("Cancelar venta")
        self.btn_cancel.clicked.connect(self._cancelar)
        self.btn_print = QPushButton("Imprimir copia")
        self.btn_print.clicked.connect(self._imprimir)
        fila.addWidget(self.btn_cancel)
        fila.addWidget(self.btn_print)
        der.addLayout(fila)
        cuerpo.addLayout(der, 52)
        root.addLayout(cuerpo)

        pie = QHBoxLayout()
        b1 = QPushButton("Desglose de artículos")
        b1.clicked.connect(self._desglose)
        b2 = QPushButton("Exportar a Excel")
        b2.clicked.connect(self._excel)
        self.lbl_suma = _lbl("Filtrado  —", 16)
        self.lbl_suma.setStyleSheet(
            "color: #0F172A; background: #ECFDF5; border: 1px solid #A7F3D0; "
            "border-radius: 10px; padding: 8px 14px; font-weight: 400; letter-spacing: 0px;"
        )
        pie.addWidget(b1)
        pie.addWidget(b2)
        pie.addWidget(self.lbl_suma)
        pie.addStretch()
        root.addLayout(pie)
        self._armado = True

    def _sin_historico(self):
        if self.chk_todo.isChecked():
            self.chk_todo.blockSignals(True)
            self.chk_todo.setChecked(False)
            self.chk_todo.blockSignals(False)

    def _ir_hoy(self):
        self._sin_historico()
        self.fecha.blockSignals(True)
        self.fecha.setDate(QDate.currentDate())
        self.fecha.blockSignals(False)
        self.refrescar()

    def _al_cambiar_dia(self):
        self._sin_historico()
        self.refrescar()

    def _pintar_dias_calendario(self):
        from src.jefe.reportes.periodo.dialogo.dias_trabajados import (
            fechas_trabajadas,
            pintar_dias_trabajados,
        )

        cal = self.fecha.calendarWidget()
        if cal is None:
            return
        pintar_dias_trabajados(cal, fechas_trabajadas())

    def refrescar(self):
        if not self._armado:
            return
        ver_todo = self.chk_todo.isChecked()
        caja_filtro = self.cb_caja.currentData()
        t0 = self.t0.time()
        t1 = self.t1.time()
        filas, total = motor_historial.listar(
            texto=self.txt.text(),
            metodo=self.cb_pago.currentText(),
            fecha_iso=self.fecha.date().toString("yyyy-MM-dd"),
            ver_todo=ver_todo,
            hora_desde=f"{t0.hour():02d}:{t0.minute():02d}",
            hora_hasta=f"{t1.hour():02d}:{t1.minute():02d}",
            solo_caja=False,
            caja_filtro=caja_filtro,
        )
        self.tabla.setRowCount(len(filas))
        for i, r in enumerate(filas):
            f_raw = r["fecha"]
            f_str = f_raw.strftime("%Y-%m-%d %H:%M:%S") if hasattr(f_raw, "strftime") else str(f_raw or "")
            dt = None
            try:
                dt = datetime.strptime(f_str[:19], "%Y-%m-%d %H:%M:%S")
            except ValueError:
                pass
            hora = dt.strftime("%H:%M") if dt else f_str[11:16]
            cancel = str(r["estado"] or "").upper().startswith("CANCELAD")
            caja = r["caja_id"] if "caja_id" in r.keys() else 1
            self.tabla.setItem(i, 0, QTableWidgetItem(str(r["id"])))
            self.tabla.setItem(i, 1, QTableWidgetItem(str(caja or 1)))
            self.tabla.setItem(i, 2, QTableWidgetItem(str(int(float(r["cant_arts"] or 0)))))
            self.tabla.setItem(i, 3, QTableWidgetItem(hora))
            tot = QTableWidgetItem(fmt_plata(r["total"]))
            tot.setForeground(QColor("#B91C1C" if cancel else "#0F172A"))
            if not cancel:
                tot.setBackground(QBrush(QColor("#F0FDF4")))
            self.tabla.setItem(i, 4, tot)
            est = QTableWidgetItem("Cancelada" if cancel else "Cerrada")
            if cancel:
                est.setForeground(QColor("#B91C1C"))
            self.tabla.setItem(i, 5, est)
        n_ok = sum(
            1
            for r in filas
            if not str(r["estado"] or "").upper().startswith("CANCELAD")
        )
        if filas:
            self.lbl_suma.setText(f"Filtrado  {fmt_plata(total)}  ·  {n_ok} tickets")
        else:
            self.lbl_suma.setText("Filtrado  $0,00")
        if self.ticket_id:
            self._pintar_detalle(self.ticket_id)

    def _detalle(self):
        items = self.tabla.selectedItems()
        if not items:
            return
        tid = int(self.tabla.item(items[0].row(), 0).text())
        self._pintar_detalle(tid)

    def _pintar_detalle(self, tid: int):
        v, lineas = motor_historial.detalle(tid)
        if not v:
            return
        self.ticket_id = v["id"]
        self.lbl_ticket.setText(f"Ticket {v['id']}")
        self.lbl_cajero.setText(f"Cajero  {str(v['usuario'] or '—')}")
        self.lbl_pago.setText(f"Pago  {str(v['metodo_pago'] or 'Efectivo')}")
        caja = v["caja_id"] if hasattr(v, "keys") and "caja_id" in v.keys() else "—"
        self.lbl_caja.setText(f"Caja  {caja}")
        self.lbl_cuando.setText(str(v["fecha"] or ""))
        audit = motor_historial.linea_cancel(v)
        if audit:
            self.lbl_audit.setText(audit)
            self.lbl_audit.show()
        else:
            self.lbl_audit.hide()
        self.tabla_det.setRowCount(len(lineas or []))
        for i, it in enumerate(lineas or []):
            self.tabla_det.setItem(i, 0, QTableWidgetItem(f"{it['cantidad']:g}"))
            self.tabla_det.setItem(i, 1, QTableWidgetItem(str(it["nombre_producto"])))
            imp = QTableWidgetItem(fmt_plata(it["subtotal"]))
            imp.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            self.tabla_det.setItem(i, 2, imp)
        self.lbl_total.setText(f"Total  {fmt_plata(v['total'])}")
        cancel = "CANCELAD" in str(v["estado"] or "").upper()
        self.btn_cancel.setEnabled(not cancel)

    def _cancelar(self):
        if not self.ticket_id:
            return
        if QMessageBox.question(
            self,
            "Cancelar venta",
            f"¿Cancelar el ticket {self.ticket_id}? Vuelve el stock y queda firmado quién cancela y desde qué caja.",
        ) != QMessageBox.StandardButton.Yes:
            return
        from src.config import config

        user = (config.current_user or {}).get("username", "jefe")
        if motor_historial.cancelar(self.ticket_id, user):
            QMessageBox.information(self, "Listo", "Ticket cancelado. Quedó la línea de auditoría.")
        else:
            QMessageBox.warning(self, "Error", "No se pudo cancelar.")
        self.refrescar()

    def _imprimir(self):
        if not self.ticket_id:
            return
        try:
            motor_historial.reimprimir(self.ticket_id)
            QMessageBox.information(self, "Listo", f"Copia del ticket {self.ticket_id} enviada.")
        except Exception as e:
            QMessageBox.critical(self, "Impresora", str(e))

    def _desglose(self):
        ids = []
        for i in range(self.tabla.rowCount()):
            it = self.tabla.item(i, 0)
            if it:
                ids.append(it.text())
        rows = motor_historial.desglose(ids)
        dlg = QDialog(self)
        dlg.setWindowTitle("Desglose de artículos")
        dlg.resize(560, 420)
        lay = QVBoxLayout(dlg)
        t = QTableWidget()
        t.setColumnCount(3)
        t.setHorizontalHeaderLabels(["Producto", "Cantidad", "Monto"])
        t.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        t.setRowCount(len(rows))
        for i, r in enumerate(rows):
            t.setItem(i, 0, QTableWidgetItem(str(r["nombre_producto"])))
            t.setItem(i, 1, QTableWidgetItem(f"{r['total_cant']:g}"))
            t.setItem(i, 2, QTableWidgetItem(fmt_plata(r["total_monto"])))
        lay.addWidget(t)
        qt_exec(dlg)

    def _excel(self):
        if self.tabla.rowCount() == 0:
            QMessageBox.information(self, "Excel", "No hay filas para exportar.")
            return
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            QMessageBox.warning(self, "Excel", "Falta openpyxl.")
            return
        path, _ = QFileDialog.getSaveFileName(self, "Guardar", "Historial_Tickets.xlsx", "Excel (*.xlsx)")
        if not path:
            return
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Historial"
        headers = ["Folio", "Caja", "Arts", "Hora", "Total", "Estado"]
        for c, h in enumerate(headers, 1):
            cell = ws.cell(1, c, h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1E40AF")
        for row in range(self.tabla.rowCount()):
            for col in range(6):
                it = self.tabla.item(row, col)
                ws.cell(row + 2, col + 1, it.text() if it else "")
        wb.save(path)
        QMessageBox.information(self, "Excel", f"Guardado en {path}")
