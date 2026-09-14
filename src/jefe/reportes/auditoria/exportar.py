"""Excel de todas las líneas del período, no solo las pintadas."""

from PyQt6.QtCore import QThread, pyqtSignal


class WorkerExportAudit(QThread):
    finished = pyqtSignal(bool, str)

    def __init__(self, path, headers, data):
        super().__init__()
        self.path = path
        self.headers = list(headers)
        self.data = [list(row) for row in data]

    def run(self):
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Auditoria"
            fill = PatternFill("solid", fgColor="1E40AF")
            font_h = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            thin = Border(
                left=Side(style="thin", color="CBD5E1"),
                right=Side(style="thin", color="CBD5E1"),
                top=Side(style="thin", color="CBD5E1"),
                bottom=Side(style="thin", color="CBD5E1"),
            )
            for i, h in enumerate(self.headers, 1):
                c = ws.cell(1, i, h)
                c.font = font_h
                c.fill = fill
                c.alignment = Alignment(horizontal="center")
                c.border = thin
            for ri, vals in enumerate(self.data, 2):
                for ci, val in enumerate(vals, 1):
                    out = val
                    if ci in (6, 8, 9):
                        try:
                            out = float(str(val).replace("$", "").replace(",", "").strip())
                        except ValueError:
                            pass
                    cell = ws.cell(ri, ci, out)
                    cell.border = thin
                    if ci in (8, 9):
                        cell.number_format = '"$"#,##0.00'
            ws.freeze_panes = "A2"
            wb.save(self.path)
            self.finished.emit(True, f"{len(self.data)} filas en {self.path}")
        except Exception as e:
            self.finished.emit(False, str(e))
